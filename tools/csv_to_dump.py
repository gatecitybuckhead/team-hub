#!/usr/bin/env python3
"""Turn a Drive `download_file_content` export into the dump format that
extract_metrics_week.py / extract_debrief.py read (sheet_text.load_rows).

WHY (2026-09-29): `read_file_content` stopped returning the full sheet. It now
returns a "Table Context" SUMMARY — ~15 sample rows with cells cut at ~60 chars,
and for the metrics sheet no cell data at all.

Two export routes:

  CSV  (exportMimeType=text/csv) — FIRST TAB ONLY. Fine for the debrief sheet
       (one tab). Wrong for metrics once a new quarter's tab is created ahead
       of time: on 2026-09-29 an empty "Q4 2026" tab went first and the CSV
       came back blank.
         python3 tools/csv_to_dump.py <saved-result.json | base64.txt> <out.json>

  XLSX (exportMimeType=application/vnd.openxmlformats-officedocument.spreadsheetml.sheet)
       — every tab; pick one by name. Use this for the metrics sheet.
         python3 tools/csv_to_dump.py --xlsx <saved-result.json> --sheet "Q3 2026" <out.json>
       Cells are rendered the way the CSV export shows them ($1,234.00, 69%)
       so the extractors see the same strings either way.

Source may be the saved tool-result JSON ({content: <base64>}) or a file
holding just the base64 string (when a small result came back inline).
"""
import sys, csv, io, json, base64, argparse, datetime


def _payload(src):
    raw = open(src, 'rb').read()
    if raw[:2] == b'PK':                       # already a raw .xlsx
        return raw
    raw = raw.decode().strip()
    try:
        b64 = json.loads(raw)['content']
    except (ValueError, KeyError, TypeError):
        b64 = raw
    return base64.b64decode(b64)


def _fmt(v, nf):
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, (datetime.datetime, datetime.date)):
        return f'{v.month}/{v.day}/{v.year}'
    if isinstance(v, (int, float)):
        nf = nf or ''
        if '%' in nf:
            return f'{v * 100:.0f}%'
        if '$' in nf:
            s = f'{abs(v):,.2f}' if '.00' in nf else f'{abs(v):,.0f}'
            return ('-' if v < 0 else '') + '$' + s
        return str(int(v)) if float(v).is_integer() else str(v)
    return str(v)


def rows_from_xlsx(data, sheet):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    if sheet not in wb.sheetnames:
        sys.exit(f'no tab {sheet!r}; tabs: {wb.sheetnames}')
    ws = wb[sheet]
    return [[_fmt(c.value, c.number_format) for c in r] for r in ws.iter_rows()]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('src')
    ap.add_argument('out')
    ap.add_argument('--xlsx', action='store_true')
    ap.add_argument('--sheet')
    a = ap.parse_args()
    data = _payload(a.src)
    if a.xlsx:
        if not a.sheet:
            sys.exit('--xlsx needs --sheet "<tab name>"')
        rows = rows_from_xlsx(data, a.sheet)
    else:
        rows = list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))

    def esc(c):
        return (c.replace('\\', '\\\\').replace('|', '\\|')
                 .replace('\r\n', ' ').replace('\n', ' '))
    lines = ['| ' + ' | '.join(esc(c) for c in r) + ' |' for r in rows]
    json.dump({'fileContent': '\n'.join(lines)}, open(a.out, 'w'))
    print(f'{a.out}: {len(rows)} rows', file=sys.stderr)


if __name__ == '__main__':
    main()
